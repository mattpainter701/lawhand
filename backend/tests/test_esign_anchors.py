"""A Word template's signing fields bind by the caption printed beside them.

A PDF template's fields carry rectangles. A Word template's cannot -- its
text reflows when values are filled -- so until now they bound to nothing and
were placed by hand on every generated copy. The template already knows the
caption around each field ("Client Signature:", "MOTHER"); generation renders
the field as a rule, finds the caption in the converted PDF, and places the
field at it. Deterministic where the detector guesses, and bound at
generation like a PDF template's.
"""

from io import BytesIO

from docx import Document
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.services.docx_templates import (
    SIGNING_RULE,
    fill_docx_template,
    word_signing_anchors,
)
from app.services.esign.anchors import WordAnchor, locate_word_signing_fields
from app.services.esign.placement import (
    ANCHOR_AMBIGUOUS,
    ANCHOR_NOT_FOUND,
    MISSING_SIGNER_ROLE,
    WORD_SOURCE_NOT_POSITIONABLE,
    template_placement_report,
    validate_placements,
)

RULE = "_" * 24


def pdf(*pages) -> bytes:
    """Each page is a list of ``(x, y, text)``."""
    buffer = BytesIO()
    doc = canvas.Canvas(buffer, pagesize=letter)
    for rows in pages:
        doc.setFont("Helvetica", 11)
        for x, y, text in rows:
            doc.drawString(x, y, text)
        doc.showPage()
    doc.save()
    return buffer.getvalue()


def sig(text, placement="after", field="client_signature", role="client", kind="signature"):
    return WordAnchor(field, kind, role, text, placement)


def locate(source, *anchors):
    return locate_word_signing_fields(source, list(anchors))


# ── Finding the caption ─────────────────────────────────────────────────────


def test_a_field_after_its_caption_lands_on_the_rule_that_follows():
    source = pdf([(72, 600, f"Client Signature: {RULE}")])
    placements, problems = locate(source, sig("Client Signature:"))

    assert problems == []
    [placed] = placements
    assert placed["page"] == 1 and placed["role"] == "client"
    assert placed["source"] == "anchored"
    x0, y0, x1, y1 = placed["rect"]
    # Over the rule, not over the caption.
    assert x0 > 150 and y0 == pytest.approx(596.0) and y1 - y0 == 28.0


def test_a_field_before_its_caption_lands_on_the_rule_ahead_of_it():
    source = pdf([(284, 600, f"{RULE}, MOTHER")])
    [placed], problems = locate(source, sig("MOTHER", "before", role="mother"))

    assert problems == []
    x0, _, x1, _ = placed["rect"]
    assert x0 == pytest.approx(284.0, abs=1.0) and x1 < 284 + 200


def test_a_field_under_its_caption_lands_on_the_rule_below():
    source = pdf([(72, 620, "Client:"), (72, 600, RULE)])
    [placed], problems = locate(source, sig("Client:", "below"))

    assert problems == []
    assert placed["rect"][1] == pytest.approx(596.0)


def test_a_caption_with_no_rule_still_gets_a_box_beside_it():
    source = pdf([(72, 600, "Client Signature:")])
    [placed], problems = locate(source, sig("Client Signature:"))

    assert problems == []
    x0, y0, x1, y1 = placed["rect"]
    assert x0 > 150 and x1 - x0 == 200.0


def test_a_missing_caption_is_reported_for_that_field_alone():
    source = pdf([(72, 600, f"Client Signature: {RULE}")])
    placements, problems = locate(
        source, sig("Client Signature:"), sig("Witness:", field="witness_signature", role="witness")
    )

    assert [item["role"] for item in placements] == ["client"]
    [problem] = problems
    assert problem.code == ANCHOR_NOT_FOUND and problem.field == "witness_signature"
    assert "'Witness:'" in problem.detail


def test_a_caption_printed_twice_is_ambiguous_unless_one_has_the_rule():
    twice = pdf([(72, 700, "Client Signature: is required below."), (72, 600, f"Client Signature: {RULE}")])
    [placed], problems = locate(twice, sig("Client Signature:"))
    assert problems == [] and placed["rect"][1] == pytest.approx(596.0)

    both = pdf([(72, 700, f"Client Signature: {RULE}"), (72, 600, f"Client Signature: {RULE}")])
    placements, problems = locate(both, sig("Client Signature:"))
    assert placements == []
    assert problems[0].code == ANCHOR_AMBIGUOUS and "2 times" in problems[0].detail


def test_the_caption_matches_across_whitespace_and_then_case():
    source = pdf([(72, 600, f"CLIENT  SIGNATURE: {RULE}")])
    [placed], problems = locate(source, sig("Client Signature:"))

    assert problems == [] and placed["page"] == 1


def test_an_exact_case_match_wins_over_a_loose_one():
    source = pdf([(72, 700, f"client signature: {RULE}"), (72, 600, f"Client Signature: {RULE}")])
    [placed], problems = locate(source, sig("Client Signature:"))

    assert problems == [] and placed["rect"][1] == pytest.approx(596.0)


def test_a_signing_date_gets_a_date_sized_box():
    source = pdf([(72, 600, f"Dated: {RULE}")])
    [placed], _ = locate(source, sig("Dated:", field="signed_on", kind="date"))

    assert placed["field_type"] == "date"
    assert placed["rect"][3] - placed["rect"][1] == 24.0


def test_fields_are_found_on_any_page():
    source = pdf([(72, 600, "Page one.")], [(72, 300, f"By: {RULE}")])
    [placed], problems = locate(source, sig("By:", role="attorney"))

    assert problems == [] and placed["page"] == 2


def test_a_field_with_no_caption_at_all_says_so():
    _, [problem] = locate(pdf([(72, 600, RULE)]), sig(""))

    assert problem.code == ANCHOR_NOT_FOUND and "no caption" in problem.detail


def test_every_anchored_box_passes_the_dispatch_validator():
    source = pdf(
        [(72, 600, f"Client Signature: {RULE}"), (72, 560, f"Dated: {RULE}"), (284, 500, f"{RULE}, MOTHER")]
    )
    placements, problems = locate(
        source,
        sig("Client Signature:"),
        sig("Dated:", field="signed_on", kind="date"),
        sig("MOTHER", "before", field="mother_signature", role="mother"),
    )

    assert problems == []
    validate_placements(
        placements,
        source_sha256=placements[0]["source_sha256"],
        signer_roles={"client", "mother"},
        required_roles=["client", "mother"],
    )


def test_unreadable_bytes_are_reported_not_raised():
    placements, problems = locate(b"not a pdf", sig("Client:"))

    assert placements == [] and problems[0].code == "unreadable_pdf"


# ── Reading the caption out of the Word document ────────────────────────────


def docx(*paragraphs) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def field(name, *, anchor=None, source_text=None, **extra):
    value = {"name": name, "field_type": "signature", "signer_role": "client", **extra}
    if anchor is not None:
        value["docx_anchor"] = anchor
    if source_text is not None:
        value["source_text"] = source_text
    return value


def test_the_caption_before_the_blank_anchors_the_field_after_it():
    source = docx("Client Signature: ____")
    [anchor] = word_signing_anchors(
        source,
        {"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": 18, "end": 22}, source_text="____")]},
    )

    assert (anchor.text, anchor.placement) == ("Client Signature:", "after")
    assert anchor.role == "client" and anchor.field_type == "signature"


def test_the_caption_after_the_blank_anchors_the_field_before_it():
    source = docx("____, MOTHER")
    [anchor] = word_signing_anchors(
        source,
        {"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": 0, "end": 4}, source_text="____", signer_role="mother")]},
    )

    assert (anchor.text, anchor.placement) == ("MOTHER", "before")


def test_a_blank_alone_on_its_line_anchors_below_the_line_above():
    source = docx("Client:", "", "____")
    [anchor] = word_signing_anchors(
        source,
        {"fields": [field("sig", anchor={"paragraph_ordinal": 2, "start": 0, "end": 4}, source_text="____")]},
    )

    assert (anchor.text, anchor.placement) == ("Client:", "below")


def test_a_placeholder_token_is_located_by_its_text():
    source = docx("By: {{attorney_signature}}")
    [anchor] = word_signing_anchors(
        source, {"fields": [field("attorney_signature", signer_role="attorney")]}
    )

    assert (anchor.text, anchor.placement) == ("By:", "after")


def test_a_sibling_placeholder_is_not_part_of_the_caption():
    # The caption is read from the unfilled template, where the value is still
    # a token. Kept, it would never match the filled PDF's "I, Ada Lovelace,".
    source = docx("I, {{client_name}}, sign here: ____")
    [anchor] = word_signing_anchors(
        source,
        {
            "fields": [
                {"name": "client_name", "field_type": "text"},
                field("sig", anchor={"paragraph_ordinal": 0, "start": 31, "end": 35}, source_text="____"),
            ]
        },
    )

    assert (anchor.text, anchor.placement) == ("sign here:", "after")


def test_a_sibling_signing_field_is_not_part_of_the_caption():
    # The signature field's own blank is rewritten as a rule when the document
    # is generated, so the date field's caption must stop short of it.
    source = docx("Signature: ____  Date: ____")
    anchors = word_signing_anchors(
        source,
        {
            "fields": [
                field("sig", anchor={"paragraph_ordinal": 0, "start": 11, "end": 15}, source_text="____"),
                field("signed_on", anchor={"paragraph_ordinal": 0, "start": 23, "end": 27}, source_text="____", field_type="date"),
            ]
        },
    )

    assert [(item.field, item.text) for item in anchors] == [
        ("sig", "Signature:"),
        ("signed_on", "Date:"),
    ]


def test_a_long_caption_is_trimmed_to_the_words_beside_the_field():
    # A caption longer than a printed line can never match one extracted line,
    # so the derivation keeps the end nearest the field.
    lead = "The undersigned, having read and understood every term of this agreement, signs below: "
    source = docx(lead + "____")
    [anchor] = word_signing_anchors(
        source,
        {"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": len(lead), "end": len(lead) + 4}, source_text="____")]},
    )

    assert len(anchor.text) <= 60
    assert anchor.text.endswith("signs below:") and lead.strip().endswith(anchor.text)


def test_a_rule_printed_beside_the_field_is_not_part_of_the_caption():
    source = docx("Name ________ Signature: ____")
    [anchor] = word_signing_anchors(
        source,
        {"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": 25, "end": 29}, source_text="____")]},
    )

    assert anchor.text == "Signature:"


def test_the_authors_anchor_wins_over_the_derived_one():
    source = docx("Client Signature: ____")
    [anchor] = word_signing_anchors(
        source,
        {
            "fields": [
                field(
                    "sig",
                    anchor={"paragraph_ordinal": 0, "start": 18, "end": 22},
                    source_text="____",
                    pdf_anchor={"text": "Signed by the client", "placement": "below"},
                )
            ]
        },
    )

    assert (anchor.text, anchor.placement) == ("Signed by the client", "below")


@pytest.mark.parametrize("override", [{"text": ""}, {"text": "   "}, "nonsense", {"placement": "below"}])
def test_an_empty_or_malformed_override_falls_back_to_the_derived_caption(override):
    source = docx("Client Signature: ____")
    [anchor] = word_signing_anchors(
        source,
        {"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": 18, "end": 22}, source_text="____", pdf_anchor=override)]},
    )

    assert anchor.text == "Client Signature:"


def test_an_unknown_placement_in_the_override_reads_as_after():
    [anchor] = word_signing_anchors(
        docx("x"), {"fields": [field("sig", pdf_anchor={"text": "By:", "placement": "sideways"})]}
    )

    assert anchor.placement == "after"


def test_a_field_whose_span_cannot_be_found_has_no_caption():
    [anchor] = word_signing_anchors(docx("Nothing here."), {"fields": [field("sig")]})

    assert anchor.text == ""


def test_only_included_signing_fields_get_anchors():
    anchors = word_signing_anchors(
        docx("Client Signature: ____", "Name: {{client_name}}"),
        {
            "fields": [
                field("sig", anchor={"paragraph_ordinal": 0, "start": 18, "end": 22}, source_text="____"),
                field("old", anchor={"paragraph_ordinal": 0, "start": 18, "end": 22}, source_text="____", included=False),
                {"name": "client_name", "field_type": "text"},
            ]
        },
    )

    assert [anchor.field for anchor in anchors] == ["sig"]


# ── The generated Word document prints a rule where the signer signs ────────


def paragraph_texts(content: bytes) -> list[str]:
    return [paragraph.text for paragraph in Document(BytesIO(content)).paragraphs]


def test_a_signing_field_prints_as_a_rule_in_the_generated_document():
    source = docx("Client Signature: [sign here]")
    filled = fill_docx_template(
        source,
        variable_schema={"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": 18, "end": 29}, source_text="[sign here]")]},
        variables={},
    )

    assert paragraph_texts(filled)[0] == f"Client Signature: {SIGNING_RULE}"


def test_a_paragraph_that_already_draws_its_rule_is_left_blank_at_the_field():
    source = docx("Client Signature: [sign here] ________")
    filled = fill_docx_template(
        source,
        variable_schema={"fields": [field("sig", anchor={"paragraph_ordinal": 0, "start": 18, "end": 29}, source_text="[sign here]")]},
        variables={},
    )

    assert paragraph_texts(filled)[0] == "Client Signature:  ________"


def test_another_fields_rule_does_not_take_this_fields_line_away():
    # The check is for a rule already drawn *beside* this field. Read across
    # the whole paragraph, a neighbour's blank left the signer nothing to
    # sign on -- and which field lost it depended on their order.
    source = docx("Name __________ Signature ________________")
    filled = fill_docx_template(
        source,
        variable_schema={
            "fields": [
                {"name": "client_name", "field_type": "text", "docx_anchor": {"paragraph_ordinal": 0, "start": 5, "end": 15}, "source_text": "__________"},
                field("sig", anchor={"paragraph_ordinal": 0, "start": 26, "end": 42}, source_text="_" * 16),
            ]
        },
        variables={"client_name": "Ada Lovelace"},
    )

    assert paragraph_texts(filled)[0] == f"Name Ada Lovelace Signature {SIGNING_RULE}"


def test_a_placeholder_token_signing_field_also_prints_as_a_rule():
    filled = fill_docx_template(
        docx("By: {{attorney_signature}}"),
        variable_schema={"fields": [field("attorney_signature", signer_role="attorney")]},
        variables={},
    )

    assert paragraph_texts(filled)[0] == f"By: {SIGNING_RULE}"


def test_a_signing_field_still_refuses_a_value():
    with pytest.raises(Exception, match="must remain blank"):
        fill_docx_template(
            docx("By: {{attorney_signature}}"),
            variable_schema={"fields": [field("attorney_signature", signer_role="attorney")]},
            variables={"attorney_signature": "Ann"},
        )


# ── Generation binds a Word template like a PDF one ─────────────────────────


def test_a_word_template_converted_to_pdf_binds_by_its_anchors():
    converted = pdf([(72, 600, f"Client Signature: {RULE}"), (72, 560, f"Dated: {RULE}")])
    report = template_placement_report(
        {"fields": [field("sig"), field("signed_on", field_type="date")]},
        source=converted,
        template_format="docx",
        output_format="pdf",
        word_anchors=[sig("Client Signature:", field="sig"), sig("Dated:", field="signed_on", kind="date")],
    )

    assert report.problems == [] and report.blocked is False
    assert [item["field_type"] for item in report.placements] == ["signature", "date"]
    assert report.recoverable is True


def test_an_anchor_that_is_not_found_blocks_that_field_and_stays_recoverable():
    converted = pdf([(72, 600, "Nothing signed here.")])
    report = template_placement_report(
        {"fields": [field("sig")]},
        source=converted,
        template_format="docx",
        output_format="pdf",
        word_anchors=[sig("Client Signature:", field="sig")],
    )

    assert report.blocked is True and report.recoverable is True
    assert [item.code for item in report.problems] == [ANCHOR_NOT_FOUND]


def test_an_anchored_field_without_a_signer_role_is_named_here_not_at_dispatch():
    converted = pdf([(72, 600, f"Client Signature: {RULE}")])
    placements, problems = locate(converted, sig("Client Signature:", role=""))

    assert placements == []
    assert [item.code for item in problems] == [MISSING_SIGNER_ROLE]
    assert "requires a signer role" in problems[0].detail


def test_without_anchors_a_word_template_reports_as_before():
    report = template_placement_report(
        {"fields": [field("sig")]},
        source=pdf([(72, 600, "x")]),
        template_format="docx",
        output_format="pdf",
    )

    assert [item.code for item in report.problems] == [WORD_SOURCE_NOT_POSITIONABLE]


def test_anchors_are_ignored_for_a_pdf_template():
    report = template_placement_report(
        {"fields": [field("sig", pdf_overlays=[{"page": 1, "rect": [72, 100, 272, 136]}])]},
        source=pdf([(72, 600, "x")]),
        template_format="pdf",
        output_format="pdf",
        word_anchors=[sig("never used")],
    )

    assert report.problems == [] and report.placements[0]["field_id"] == "field-0-0"
