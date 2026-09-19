"""Mock documents for the fill campaign, built at test time.

The AcroForm PDF and the Word document carry the same field names, chosen to
cover every naming convention a firm's own form might use: the alias Smart
Fill hardcodes, its spaced and camel-cased spellings, the natural names a
person types (``first_name``, ``Client Name``), the caption roles the party
model accepts, and the matter columns nobody can reach.  A ``/Sig`` widget,
a choice with long-form options, and a checkbox exercise the renderer's
non-text paths.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DictionaryObject,
    FloatObject,
    NameObject,
    NumberObject,
    TextStringObject,
)
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.services.pdf_templates import discover_pdf_fields

SEED_ROOT = Path(__file__).resolve().parents[2] / "seed" / "sample_templates"

#: PDF field names, in the order they are drawn. The comment says which
#: convention each one probes.
TEXT_FIELDS: tuple[str, ...] = (
    "client_name",  # the hardcoded alias
    "Client Name",  # spaced: folds to client_name, collides, gets a suffix
    "ClientName",  # camel-cased: folds to clientname, matches nothing
    "client_full_name",  # the most common customer spelling
    "first_name",  # contact.first_name, unreachable
    "last_name",  # contact.last_name, unreachable
    "client_email",
    "client_city",
    "petitioner_name",  # family-law caption role
    "respondent_name",
    "plaintiff_name",  # civil caption role
    "defendant_name",
    "defendant_2_full_name",  # instance alias for a second defendant
    "matter_number",  # unreachable matter column
    "case_number",
    "court",
    "opened_on",  # unreachable date column
    "hourly_rate",  # unformatted Decimal
    "retainer_amount",
    "estate_decedent_name",  # loads only for an estate.* binding
    "prepared_by",
    "attorney_name",
    "firm_name",  # resolves only through a firm.* binding
)

CHOICE_FIELD = "client_state"
CHOICE_OPTIONS = ("North Dakota", "Minnesota", "South Dakota")
CHECKBOX_FIELD = "agree"
SIGNATURE_FIELD = "client_signature"


def convention_pdf() -> bytes:
    """An AcroForm with every convention above, one page."""

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setFont("Helvetica", 8)
    y = 760
    for name in TEXT_FIELDS:
        pdf.drawString(36, y + 4, name)
        pdf.acroForm.textfield(
            name=name, tooltip=name, x=200, y=y, width=260, height=14, borderWidth=0
        )
        y -= 20
    pdf.drawString(36, y + 4, CHOICE_FIELD)
    pdf.acroForm.choice(
        name=CHOICE_FIELD,
        x=200,
        y=y,
        width=160,
        height=14,
        options=list(CHOICE_OPTIONS),
        value=CHOICE_OPTIONS[0],
    )
    y -= 20
    pdf.drawString(36, y + 4, CHECKBOX_FIELD)
    pdf.acroForm.checkbox(name=CHECKBOX_FIELD, x=200, y=y, size=12, fieldFlags="")
    y -= 40
    pdf.drawString(36, y + 4, "Client signature")
    pdf.showPage()
    pdf.save()

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(buffer.getvalue())))
    page = writer.pages[0]
    signature = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/FT"): NameObject("/Sig"),
            NameObject("/T"): TextStringObject(SIGNATURE_FIELD),
            NameObject("/TU"): TextStringObject("Client signature"),
            NameObject("/Rect"): ArrayObject(
                [
                    FloatObject(200),
                    FloatObject(y - 4),
                    FloatObject(460),
                    FloatObject(y + 20),
                ]
            ),
            NameObject("/F"): NumberObject(4),
            NameObject("/P"): page.indirect_reference,
        }
    )
    reference = writer._add_object(signature)
    page[NameObject("/Annots")].append(reference)
    writer._root_object["/AcroForm"]["/Fields"].append(reference)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def convention_docx() -> tuple[bytes, list[str]]:
    """A Word letter using ``{{name}}`` placeholders for the same conventions."""

    from docx import Document

    names = [name for name in TEXT_FIELDS if name != "ClientName"] + ["ClientName"]
    document = Document()
    document.add_paragraph("Engagement letter")
    document.add_paragraph("Dear {{first_name}} {{last_name}},")
    document.add_paragraph(
        "Client: {{client_name}} / {{Client Name}} / {{ClientName}} / {{client_full_name}}"
    )
    document.add_paragraph("Contact: {{client_email}}, {{client_city}}")
    document.add_paragraph("Caption: {{petitioner_name}} v. {{respondent_name}}")
    document.add_paragraph(
        "Caption: {{plaintiff_name}} v. {{defendant_name}} and {{defendant_2_full_name}}"
    )
    document.add_paragraph(
        "Matter {{matter_number}}, case {{case_number}}, {{court}}, opened {{opened_on}}"
    )
    document.add_paragraph(
        "Rate {{hourly_rate}} per hour; retainer {{retainer_amount}}"
    )
    document.add_paragraph("Estate of {{estate_decedent_name}}")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Prepared by"
    table.cell(0, 1).text = "{{prepared_by}}"
    table.cell(1, 0).text = "Attorney"
    table.cell(1, 1).text = "{{attorney_name}} of {{firm_name}}"
    document.add_paragraph("Signed: {{client_signature}}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), names + [SIGNATURE_FIELD]


def schema_for_pdf(
    pdf: bytes,
    *,
    bindings: dict[str, str] | None = None,
    manual: tuple[str, ...] = (),
) -> dict[str, Any]:
    """Discover the PDF's fields the way template creation does, then bind some.

    ``bindings`` is keyed by the *discovered variable name*. The renderer
    demands a schema row for every widget the PDF carries, so every discovered
    field is included whether or not a binding is attached.
    """

    fields = []
    for discovered in discover_pdf_fields(pdf):
        row = dict(discovered)
        row["included"] = True
        name = row["name"]
        if bindings and name in bindings:
            row["binding"] = bindings[name]
        if name in manual:
            row["binding"] = "manual"
        fields.append(row)
    return {"fields": fields}


def schema_for_docx(
    names: list[str],
    *,
    bindings: dict[str, str] | None = None,
    manual: tuple[str, ...] = (),
) -> dict[str, Any]:
    fields = []
    for name in names:
        row: dict[str, Any] = {"name": name, "type": "text", "included": True}
        if name == SIGNATURE_FIELD:
            row["type"] = "signature"
            row["field_type"] = "signature"
        if bindings and name in bindings:
            row["binding"] = bindings[name]
        if name in manual:
            row["binding"] = "manual"
        fields.append(row)
    return {"fields": fields}


#: Bindings a customer would declare on the convention form once, in the
#: editor. Left out on purpose: ``first_name``/``last_name`` (no catalogue
#: path exists) and ``matter_number``/``opened_on`` (same).
CONVENTION_BINDINGS: dict[str, str] = {
    "client_full_name": "client.name",
    "clientname": "client.name",
    "client_name_2": "client.name",
    "client_email": "client.email",
    "client_city": "client.address.city",
    "client_state": "client.address.state",
    "plaintiff_name": "plaintiff.full_name",
    "defendant_name": "defendant.full_name",
    "defendant_2_full_name": "defendant.2.full_name",
    "case_number": "matter.case_number",
    "court": "matter.court",
    "hourly_rate": "matter.hourly_rate",
    "retainer_amount": "matter.retainer_amount",
    "estate_decedent_name": "estate.decedent_name",
    "attorney_name": "attorney.name",
    "firm_name": "firm.name",
}


#: The Word document keeps the spaced and camel-cased spellings as written,
#: so its binding map is keyed by those names rather than the PDF's folded
#: and suffixed variants.
DOCX_BINDINGS: dict[str, str] = {
    **{
        key: value
        for key, value in CONVENTION_BINDINGS.items()
        if key not in ("clientname", "client_name_2")
    },
    "ClientName": "client.name",
    "Client Name": "client.name",
}


def manifest() -> list[dict[str, Any]]:
    return json.loads((SEED_ROOT / "manifest.json").read_text())["forms"]


def seed_forms_with_bindings() -> list[dict[str, Any]]:
    return [form for form in manifest() if form.get("bindings")]


def seed_pdf(form: dict[str, Any]) -> bytes:
    return (SEED_ROOT / form["filename"]).read_bytes()
