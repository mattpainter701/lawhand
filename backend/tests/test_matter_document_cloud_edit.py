"""Open a matter's Word document in the firm's Word or Google Docs, and bring edits back.

These drive the real ASGI app against Postgres with the provider calls faked
at the store boundary. They pin the product rules: the document keeps its
identity when edits come back, approved and signing-bound documents are
never overwritten, an upload never lands on top of a newer Word edit, and
every adoption leaves an integrity and timeline record.
"""

import hashlib
from datetime import datetime, timezone
import io
import uuid
from types import SimpleNamespace

import pytest
import pytest_asyncio
from docx import Document
from sqlalchemy import select

from app.models.document_integrity_event import DocumentIntegrityEvent
from app.models.matter_document import MatterDocument
from app.models.plugin import Matter, MatterEvent
from app.models.signature import SignatureRequest
from app.routers import matter_documents as routes
from app.services.matter_file_store import (
    MatterFileIntegrityError,
    MatterFileNotFound,
    ProviderFileMetadata,
    StorageResult,
)
from app.services.provider_http import ProviderAuthError, ProviderError

API = "/api/matters"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


ORIGINAL = _docx("Engagement letter, as generated")
EDITED = _docx("Engagement letter, edited in Word")


def _sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


class FakeCloud:
    """The firm's cloud file as the store sees it."""

    def __init__(self, content: bytes, *, backend: str = "sharepoint") -> None:
        self.content = content
        self.backend = backend
        self.etag = '"{E},1"'
        self.read_error: Exception | None = None
        self.metadata_error: Exception | None = None
        self.replace_error: Exception | None = None
        self.replaced: list[dict] = []

    def install(self, monkeypatch) -> None:
        store = routes.matter_file_store

        async def read(**kwargs):
            if self.read_error:
                raise self.read_error
            assert kwargs["expected_sha256"] is None
            assert kwargs["enforce_persisted_size"] is False
            return self.content

        async def metadata(**kwargs):
            if self.metadata_error:
                raise self.metadata_error
            if self.backend == "google_drive":
                return ProviderFileMetadata(
                    backend="google_drive",
                    web_url="https://docs.google.com/document/d/item-1/edit",
                    version_id="7",
                )
            return ProviderFileMetadata(
                backend=self.backend,
                web_url="https://firm.sharepoint.com/:w:/r/Doc.aspx?sourcedoc=1",
                desktop_url="https://firm.sharepoint.com/sites/m/Shared%20Documents/a.docx",
                etag=self.etag,
                version_id='"c:{E},1"',
                modified_at="2026-09-25T10:00:00Z",
            )

        async def replace(**kwargs):
            if self.replace_error:
                raise self.replace_error
            self.replaced.append(kwargs)
            self.content = kwargs["content"]
            self.etag = '"{E},9"'
            return StorageResult(
                provider="microsoft",
                backend=self.backend,
                provider_item_id="item-1",
                provider_etag=self.etag,
                provider_version_id='"c:{E},9"',
                provider_modified_at="2026-09-25T11:00:00Z",
            )

        monkeypatch.setattr(store, "read_matter_file_bytes", read)
        monkeypatch.setattr(store, "get_matter_file_metadata", metadata)
        monkeypatch.setattr(store, "replace_matter_file_content", replace)


@pytest_asyncio.fixture(autouse=True)
async def _no_index_jobs(monkeypatch):
    queued = []

    async def enqueue_index(**kwargs):
        queued.append(kwargs)
        return True

    monkeypatch.setattr(routes.matter_document_index, "enqueue_index", enqueue_index)
    return queued


@pytest_asyncio.fixture
async def matter(db_session, test_tenant, test_user):
    row = Matter(
        id=uuid.uuid4(),
        tenant_id=test_tenant.id,
        user_id=test_user.id,
        slug=f"cloud-edit-{uuid.uuid4().hex[:8]}",
        matter_name="Atlas acquisition",
        matter_type="corporate",
        status="open",
    )
    db_session.add(row)
    await db_session.commit()
    return SimpleNamespace(id=row.id, tenant_id=test_tenant.id, user_id=test_user.id)


async def _document(db_session, matter, **overrides) -> uuid.UUID:
    fields = dict(
        id=uuid.uuid4(),
        tenant_id=matter.tenant_id,
        matter_id=matter.id,
        uploaded_by_user_id=matter.user_id,
        filename="Engagement letter.docx",
        content_type=DOCX,
        file_size=len(ORIGINAL),
        storage_path="https://firm.sharepoint.com/display",
        storage_provider="microsoft",
        storage_backend="sharepoint",
        provider_object_id="item-1",
        provider_drive_id="drive-1",
        provider_etag='"{E},1"',
        document_sha256=_sha(ORIGINAL),
        document_status="draft",
    )
    fields.update(overrides)
    row = MatterDocument(**fields)
    db_session.add(row)
    await db_session.commit()
    return row.id


async def _load(db_session, doc_id) -> MatterDocument:
    db_session.expire_all()
    return (
        await db_session.execute(select(MatterDocument).where(MatterDocument.id == doc_id))
    ).scalar_one()


def _url(matter, doc_id, action):
    return f"{API}/{matter.id}/documents/{doc_id}/{action}"


# --- Open in Word / Google Docs -------------------------------------------


@pytest.mark.asyncio
async def test_open_returns_word_links_and_records_who_is_editing(client, db_session, matter, monkeypatch):
    FakeCloud(ORIGINAL).install(monkeypatch)
    doc_id = await _document(db_session, matter)

    resp = await client.post(_url(matter, doc_id, "cloud-edit"), json={"app": "word_desktop"})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["links"] == {
        "word_web": "https://firm.sharepoint.com/:w:/r/Doc.aspx?sourcedoc=1",
        "word_desktop": "ms-word:ofe|u|https://firm.sharepoint.com/sites/m/Shared%20Documents/a.docx",
    }
    assert body["app"] == "word_desktop"
    assert body["document"]["external_edit_app"] == "word_desktop"
    assert body["document"]["external_edit_started_by_name"] == "Test Attorney"

    listing = await client.get(f"{API}/{matter.id}/documents")
    row = next(item for item in listing.json()["items"] if item["id"] == str(doc_id))
    assert row["external_edit_started_at"]
    assert row["external_edit_started_by_name"] == "Test Attorney"


@pytest.mark.asyncio
async def test_open_defaults_to_google_docs_for_a_drive_file(client, db_session, matter, monkeypatch):
    FakeCloud(ORIGINAL, backend="google_drive").install(monkeypatch)
    doc_id = await _document(
        db_session, matter, storage_backend="google_drive", storage_provider="google", provider_drive_id=None
    )
    resp = await client.post(_url(matter, doc_id, "cloud-edit"), json={})
    assert resp.status_code == 200, resp.text
    assert resp.json()["links"] == {"google_docs": "https://docs.google.com/document/d/item-1/edit"}
    assert resp.json()["app"] == "google_docs"

    wrong_app = await client.post(_url(matter, doc_id, "cloud-edit"), json={"app": "word_desktop"})
    assert wrong_app.status_code == 422
    assert wrong_app.json()["detail"]["code"] == "app_unavailable"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("overrides", "status", "code"),
    [
        ({"filename": "Scan.pdf", "content_type": "application/pdf"}, 422, "not_word_document"),
        ({"storage_backend": "local", "storage_provider": "local", "storage_path": "/tmp/a.docx"}, 409, "not_in_cloud"),
        ({"provider_object_id": None}, 409, "cloud_binding_missing"),
        ({"document_status": "approved"}, 409, "document_locked"),
    ],
)
async def test_open_refuses_documents_it_cannot_round_trip(
    client, db_session, matter, monkeypatch, overrides, status, code
):
    FakeCloud(ORIGINAL).install(monkeypatch)
    doc_id = await _document(db_session, matter, **overrides)
    resp = await client.post(_url(matter, doc_id, "cloud-edit"), json={})
    assert resp.status_code == status, resp.text
    assert resp.json()["detail"]["code"] == code
    assert (await _load(db_session, doc_id)).external_edit_started_at is None


@pytest.mark.asyncio
async def test_open_reports_a_missing_connection_in_plain_words(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.metadata_error = ProviderAuthError("expired")
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "cloud-edit"), json={})
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "reconnect_required"


@pytest.mark.asyncio
async def test_open_refuses_a_document_out_for_signature(client, db_session, matter, monkeypatch):
    FakeCloud(ORIGINAL).install(monkeypatch)
    doc_id = await _document(db_session, matter)
    db_session.add(
        SignatureRequest(
            tenant_id=matter.tenant_id, matter_id=matter.id, document_id=doc_id, status="sent"
        )
    )
    await db_session.commit()
    resp = await client.post(_url(matter, doc_id, "cloud-edit"), json={})
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "document_locked"


# --- Bring back changes ----------------------------------------------------


@pytest.mark.asyncio
async def test_bring_back_unchanged_verifies_and_keeps_the_editing_marker(client, db_session, matter, monkeypatch, _no_index_jobs):
    FakeCloud(ORIGINAL).install(monkeypatch)
    doc_id = await _document(db_session, matter)
    await client.post(_url(matter, doc_id, "cloud-edit"), json={})

    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.status_code == 200, resp.text
    assert resp.json()["outcome"] == "unchanged"
    doc = await _load(db_session, doc_id)
    assert doc.storage_state == "verified"
    assert doc.document_sha256 == _sha(ORIGINAL)
    # The person is usually still editing; later edits must keep coming back.
    assert doc.external_edit_started_at is not None
    assert _no_index_jobs == []


@pytest.mark.asyncio
async def test_bring_back_clears_an_editing_marker_after_twelve_hours(client, db_session, matter, monkeypatch):
    from datetime import datetime, timedelta, timezone

    FakeCloud(ORIGINAL).install(monkeypatch)
    doc_id = await _document(
        db_session,
        matter,
        external_edit_started_at=datetime.now(timezone.utc) - timedelta(hours=13),
        external_edit_started_by=matter.user_id,
        external_edit_app="word_web",
    )
    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.json()["outcome"] == "unchanged"
    doc = await _load(db_session, doc_id)
    assert doc.external_edit_started_at is None
    assert doc.external_edit_app is None


@pytest.mark.asyncio
async def test_bring_back_adopts_word_edits_in_place(client, db_session, matter, monkeypatch, _no_index_jobs):
    cloud = FakeCloud(ORIGINAL)
    cloud.install(monkeypatch)
    doc_id = await _document(
        db_session, matter, storage_state="conflict", storage_error="Download verification failed"
    )
    await client.post(_url(matter, doc_id, "cloud-edit"), json={"app": "word_web"})
    cloud.content = EDITED
    cloud.etag = '"{E},2"'

    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.status_code == 200, resp.text
    assert resp.json()["outcome"] == "adopted"
    doc = await _load(db_session, doc_id)
    assert doc.id == doc_id
    assert doc.provider_object_id == "item-1"
    assert doc.document_sha256 == _sha(EDITED)
    assert doc.file_size == len(EDITED)
    assert doc.provider_etag == '"{E},2"'
    assert doc.provider_modified_at is not None
    assert doc.storage_state == "verified"
    assert doc.storage_error is None
    assert doc.document_status == "draft"
    assert doc.external_edit_started_at is not None
    assert _no_index_jobs and _no_index_jobs[0]["document_sha256"] == _sha(EDITED)

    events = (
        await db_session.execute(
            select(DocumentIntegrityEvent).where(DocumentIntegrityEvent.document_id == doc_id)
        )
    ).scalars().all()
    adopted = [event for event in events if event.event_type == "external_cloud_edit_adopted"]
    assert len(adopted) == 1
    assert adopted[0].content_sha256 == _sha(EDITED)
    timeline = (
        await db_session.execute(
            select(MatterEvent).where(MatterEvent.event_type == "document_edits_adopted")
        )
    ).scalars().all()
    assert len(timeline) == 1
    assert timeline[0].metadata_json["previous_sha256"] == _sha(ORIGINAL)
    assert timeline[0].metadata_json["editor_app"] == "word_web"


@pytest.mark.asyncio
async def test_bring_back_never_overwrites_an_approved_document(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(EDITED)
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter, document_status="approved")

    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.status_code == 200, resp.text
    assert resp.json()["outcome"] == "blocked"
    assert "approved" in resp.json()["message"]
    doc = await _load(db_session, doc_id)
    assert doc.document_sha256 == _sha(ORIGINAL)
    assert doc.storage_state == "conflict"
    assert "approved" in doc.storage_error
    blocked = (
        await db_session.execute(
            select(DocumentIntegrityEvent).where(
                DocumentIntegrityEvent.document_id == doc_id,
                DocumentIntegrityEvent.event_type == "external_cloud_edit_blocked",
            )
        )
    ).scalars().all()
    assert len(blocked) == 1


@pytest.mark.asyncio
async def test_bring_back_records_a_baseline_for_a_legacy_row(client, db_session, matter, monkeypatch):
    FakeCloud(ORIGINAL).install(monkeypatch)
    doc_id = await _document(db_session, matter, document_sha256=None)
    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.json()["outcome"] == "unchanged"
    assert (await _load(db_session, doc_id)).document_sha256 == _sha(ORIGINAL)


@pytest.mark.asyncio
async def test_bring_back_rejects_an_unsafe_file(client, db_session, matter, monkeypatch):
    FakeCloud(b"not a word file").install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"].startswith("docx_")
    assert (await _load(db_session, doc_id)).document_sha256 == _sha(ORIGINAL)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("error", "status", "code"),
    [
        (MatterFileNotFound("gone"), 404, "cloud_file_missing"),
        (ProviderAuthError("expired"), 409, "reconnect_required"),
        (ProviderError("boom"), 503, "cloud_unavailable"),
    ],
)
async def test_bring_back_explains_provider_failures(client, db_session, matter, monkeypatch, error, status, code):
    cloud = FakeCloud(ORIGINAL)
    cloud.read_error = error
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.status_code == status
    assert resp.json()["detail"]["code"] == code


@pytest.mark.asyncio
async def test_bring_back_still_adopts_when_metadata_is_unavailable(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(EDITED)
    cloud.metadata_error = ProviderError("metadata down")
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "reconcile"))
    assert resp.json()["outcome"] == "adopted"
    doc = await _load(db_session, doc_id)
    assert doc.document_sha256 == _sha(EDITED)
    assert doc.provider_etag == '"{E},1"'


@pytest.mark.asyncio
async def test_bring_back_leaves_assistant_revisions_to_their_own_flow(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(EDITED)
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter, document_category="assistant_revision")
    for action in ("cloud-edit", "reconcile"):
        resp = await client.post(_url(matter, doc_id, action), json={})
        assert resp.status_code == 409
        assert resp.json()["detail"]["code"] == "assistant_draft"
    upload = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert upload.status_code == 409
    assert (await _load(db_session, doc_id)).document_sha256 == _sha(ORIGINAL)


def test_artifact_bound_drafts_are_refused_before_any_provider_call():
    doc = MatterDocument(
        filename="Draft.docx",
        content_type=DOCX,
        storage_backend="onedrive",
        provider_object_id="item-1",
        generated_artifact_revision_id=uuid.uuid4(),
    )
    with pytest.raises(routes.cloud_edit.CloudEditError) as refused:
        routes.cloud_edit.assert_editable_in_office(doc, require_cloud=True)
    assert refused.value.code == "assistant_draft"


# --- Upload revised version ------------------------------------------------


def _upload(content: bytes, name: str = "Engagement letter (revised).docx"):
    return {"file": (name, content, DOCX)}


@pytest.mark.asyncio
async def test_upload_writes_the_same_cloud_item_with_if_match_and_adopts(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)

    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert resp.status_code == 200, resp.text
    assert resp.json()["outcome"] == "adopted"
    assert len(cloud.replaced) == 1
    # The write is guarded by the eTag of the bytes just checked.
    assert cloud.replaced[0]["if_match"] == '"{E},1"'
    assert cloud.replaced[0]["content"] == EDITED
    doc = await _load(db_session, doc_id)
    assert doc.filename == "Engagement letter.docx"
    assert doc.document_sha256 == _sha(EDITED)
    assert doc.provider_etag == '"{E},9"'
    timeline = (
        await db_session.execute(
            select(MatterEvent).where(MatterEvent.event_type == "document_edits_adopted")
        )
    ).scalars().one()
    assert timeline.metadata_json["source"] == "revised_upload"


@pytest.mark.asyncio
async def test_upload_never_lands_on_top_of_newer_word_edits(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(_docx("Someone else edited this in Word"))
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)

    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "changed_in_office"
    assert cloud.replaced == []
    assert (await _load(db_session, doc_id)).document_sha256 == _sha(ORIGINAL)


@pytest.mark.asyncio
async def test_upload_reports_a_word_edit_that_raced_the_write(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.replace_error = MatterFileIntegrityError("412")
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "changed_in_office"
    assert (await _load(db_session, doc_id)).document_sha256 == _sha(ORIGINAL)


@pytest.mark.asyncio
async def test_upload_of_the_same_file_changes_nothing(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(ORIGINAL))
    assert resp.status_code == 200
    assert resp.json()["outcome"] == "unchanged"
    assert cloud.replaced == []


@pytest.mark.asyncio
async def test_upload_refuses_unsafe_files_and_locked_documents(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    bad = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(b"PK not really"))
    assert bad.status_code == 422
    locked_id = await _document(db_session, matter, id=uuid.uuid4(), document_status="filed")
    locked = await client.post(_url(matter, locked_id, "revised-version"), files=_upload(EDITED))
    assert locked.status_code == 409
    assert locked.json()["detail"]["code"] == "document_locked"
    assert cloud.replaced == []


@pytest.mark.asyncio
async def test_upload_works_for_a_locally_stored_document(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL, backend="local")
    cloud.install(monkeypatch)
    doc_id = await _document(
        db_session,
        matter,
        storage_backend="local",
        storage_provider="local",
        storage_path="/uploads/a.docx",
        provider_object_id=None,
        provider_drive_id=None,
        provider_etag=None,
    )
    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert resp.status_code == 200, resp.text
    assert resp.json()["outcome"] == "adopted"
    assert cloud.replaced[0]["if_match"] is None
    doc = await _load(db_session, doc_id)
    assert doc.document_sha256 == _sha(EDITED)
    assert doc.provider_etag is None


@pytest.mark.asyncio
async def test_edit_routes_are_tenant_fenced(client, db_session, matter, monkeypatch):
    FakeCloud(ORIGINAL).install(monkeypatch)
    missing = uuid.uuid4()
    for action in ("cloud-edit", "reconcile"):
        resp = await client.post(_url(matter, missing, action), json={})
        assert resp.status_code == 404


@pytest.mark.asyncio
async def test_upload_guards_with_the_fresh_etag_not_the_stored_one(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.etag = '"{E},5"'
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter, provider_etag='"{E},stale"')
    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert resp.status_code == 200, resp.text
    assert cloud.replaced[0]["if_match"] == '"{E},5"'
    doc = await _load(db_session, doc_id)
    assert doc.external_edit_started_at is None


@pytest.mark.asyncio
async def test_upload_still_writes_when_metadata_is_unavailable(client, db_session, matter, monkeypatch):
    cloud = FakeCloud(ORIGINAL)
    cloud.metadata_error = ProviderError("metadata down")
    cloud.install(monkeypatch)
    doc_id = await _document(db_session, matter)
    resp = await client.post(_url(matter, doc_id, "revised-version"), files=_upload(EDITED))
    assert resp.status_code == 200, resp.text
    assert cloud.replaced[0]["if_match"] is None


@pytest.mark.asyncio
async def test_client_portal_users_cannot_open_bring_back_or_replace(client, db_session, matter, monkeypatch):
    from datetime import timedelta

    from jose import jwt as jose_jwt

    from app.config import get_settings
    from app.models.user import User

    cloud = FakeCloud(ORIGINAL)
    cloud.install(monkeypatch)
    reads = []
    original_read = routes.matter_file_store.read_matter_file_bytes

    async def counting_read(**kwargs):
        reads.append(kwargs)
        return await original_read(**kwargs)

    monkeypatch.setattr(routes.matter_file_store, "read_matter_file_bytes", counting_read)
    doc_id = await _document(db_session, matter)
    portal_user = User(
        id=uuid.uuid4(),
        tenant_id=matter.tenant_id,
        email="client@example.test",
        full_name="Portal Client",
        role="client",
        is_active=True,
    )
    db_session.add(portal_user)
    await db_session.commit()
    settings = get_settings()
    token = jose_jwt.encode(
        {
            "sub": str(portal_user.id),
            "tenant_id": str(matter.tenant_id),
            "role": "client",
            "email": "client@example.test",
            "exp": datetime.now(timezone.utc) + timedelta(hours=1),
        },
        settings.SECRET_KEY,
        algorithm=settings.ALGORITHM,
    )
    headers = {"Authorization": f"Bearer {token}"}

    for action in ("cloud-edit", "reconcile"):
        resp = await client.post(_url(matter, doc_id, action), json={}, headers=headers)
        assert resp.status_code == 403, resp.text
        assert resp.json()["detail"]["code"] == "staff_only"
    upload = await client.post(
        _url(matter, doc_id, "revised-version"), files=_upload(EDITED), headers=headers
    )
    assert upload.status_code == 403
    assert upload.json()["detail"]["code"] == "staff_only"
    assert reads == []
    assert cloud.replaced == []
    doc = await _load(db_session, doc_id)
    assert doc.document_sha256 == _sha(ORIGINAL)
    assert doc.external_edit_started_at is None
